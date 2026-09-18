import os
import subprocess
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent

# ==============================================================================
# 1. HTML GENERATOR WITH PRINT-OPTIMIZED CSS
# ==============================================================================

HTML_CONTENT = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>ResearchMind AI — Comprehensive Project Documentation & Technical Architecture Report</title>
<style>
  @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500;600&display=swap');

  :root {
    --primary: #4F46E5;
    --primary-dark: #3730A3;
    --secondary: #7C3AED;
    --accent-orange: #F97316;
    --navy-bg: #060B18;
    --card-bg: #0D1527;
    --text-main: #1E293B;
    --text-muted: #64748B;
    --border: #E2E8F0;
    --bg-light: #F8FAFC;
    --code-bg: #F1F5F9;
  }

  @page {
    size: A4;
    margin: 20mm 18mm 20mm 18mm;
    @bottom-right {
      content: counter(page);
      font-family: 'Plus Jakarta Sans', sans-serif;
      font-size: 9pt;
      color: #94A3B8;
    }
  }

  body {
    font-family: 'Plus Jakarta Sans', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
    color: var(--text-main);
    line-height: 1.65;
    background: #FFFFFF;
    margin: 0;
    padding: 30px 40px;
    font-size: 10.5pt;
  }

  /* Cover / Header */
  .doc-header {
    border-bottom: 2px solid var(--border);
    padding-bottom: 25px;
    margin-bottom: 35px;
  }
  .brand-badge {
    display: inline-flex;
    align-items: center;
    gap: 8px;
    background: #EEF2FF;
    color: var(--primary);
    padding: 6px 14px;
    border-radius: 9999px;
    font-size: 9pt;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.05em;
    border: 1px solid #C7D2FE;
    margin-bottom: 12px;
  }
  h1.doc-title {
    font-size: 24pt;
    font-weight: 800;
    color: #0F172A;
    line-height: 1.25;
    margin: 0 0 10px 0;
    letter-spacing: -0.02em;
  }
  .doc-subtitle {
    font-size: 12pt;
    color: var(--text-muted);
    margin: 0 0 18px 0;
  }
  .meta-grid {
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    gap: 12px;
    background: var(--bg-light);
    border: 1px solid var(--border);
    border-radius: 12px;
    padding: 12px 18px;
    font-size: 9pt;
  }
  .meta-item strong {
    display: block;
    color: #475569;
    font-size: 8pt;
    text-transform: uppercase;
    letter-spacing: 0.04em;
    margin-bottom: 2px;
  }
  .meta-item span {
    font-weight: 600;
    color: #0F172A;
  }

  /* Typography */
  h2 {
    font-size: 15pt;
    font-weight: 700;
    color: #0F172A;
    border-bottom: 1.5px solid var(--border);
    padding-bottom: 6px;
    margin-top: 36px;
    margin-bottom: 14px;
    page-break-after: avoid;
  }
  h3 {
    font-size: 12pt;
    font-weight: 700;
    color: #1E293B;
    margin-top: 22px;
    margin-bottom: 8px;
    page-break-after: avoid;
  }
  h4 {
    font-size: 11pt;
    font-weight: 600;
    color: var(--primary);
    margin-top: 14px;
    margin-bottom: 6px;
  }
  p {
    margin: 0 0 12px 0;
  }
  ul, ol {
    margin: 0 0 14px 0;
    padding-left: 22px;
  }
  li {
    margin-bottom: 4px;
  }

  /* Cards & Callouts */
  .callout {
    background: #F8FAFC;
    border-left: 4px solid var(--primary);
    border-radius: 0 8px 8px 0;
    padding: 12px 16px;
    margin: 16px 0;
    font-size: 10pt;
  }
  .callout-highlight {
    background: #FEF3C7;
    border-left-color: #D97706;
  }
  .callout-title {
    font-weight: 700;
    color: #0F172A;
    margin-bottom: 4px;
  }

  /* Tables */
  table {
    width: 100%;
    border-collapse: collapse;
    margin: 16px 0 24px 0;
    font-size: 9.5pt;
    page-break-inside: avoid;
  }
  th {
    background: #0F172A;
    color: #FFFFFF;
    text-align: left;
    padding: 10px 12px;
    font-weight: 600;
    font-size: 9pt;
    letter-spacing: 0.02em;
  }
  td {
    padding: 9px 12px;
    border-bottom: 1px solid var(--border);
    vertical-align: top;
  }
  tr:nth-child(even) td {
    background: #F8FAFC;
  }

  /* Code & Pre */
  code {
    font-family: 'JetBrains Mono', Consolas, Monaco, monospace;
    font-size: 9pt;
    background: var(--code-bg);
    padding: 2px 6px;
    border-radius: 4px;
    color: #B91C1C;
  }
  pre {
    background: #0F172A;
    color: #F8FAFC;
    padding: 14px 18px;
    border-radius: 8px;
    overflow-x: auto;
    font-family: 'JetBrains Mono', Consolas, Monaco, monospace;
    font-size: 8.5pt;
    line-height: 1.5;
    margin: 14px 0;
    page-break-inside: avoid;
  }
  pre code {
    background: transparent;
    color: inherit;
    padding: 0;
  }

  /* Diagrams / ASCII Boxes */
  .diagram-box {
    background: #0D1527;
    color: #38BDF8;
    border-radius: 8px;
    padding: 14px 18px;
    font-family: 'JetBrains Mono', monospace;
    font-size: 8pt;
    line-height: 1.35;
    margin: 16px 0;
    overflow-x: auto;
    page-break-inside: avoid;
  }

  /* Pillar Grid */
  .pillar-grid {
    display: grid;
    grid-template-columns: repeat(2, 1fr);
    gap: 14px;
    margin: 16px 0;
    page-break-inside: avoid;
  }
  .pillar-card {
    border: 1px solid var(--border);
    border-radius: 10px;
    padding: 14px 16px;
    background: #FFFFFF;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04);
  }
  .pillar-card h4 {
    margin: 0 0 6px 0;
    font-size: 10.5pt;
    color: #1E293B;
  }
  .pillar-card p {
    margin: 0;
    font-size: 9pt;
    color: #475569;
  }

  /* Footer */
  .doc-footer {
    margin-top: 50px;
    border-top: 1px solid var(--border);
    padding-top: 15px;
    text-align: center;
    font-size: 8.5pt;
    color: #94A3B8;
  }

  @media print {
    body {
      padding: 0;
    }
    .page-break {
      page-break-before: always;
    }
  }
</style>
</head>
<body>

<div class="doc-header">
  <div class="brand-badge">🧠🔬 Official System Architecture &amp; Technical Report</div>
  <h1 class="doc-title">ResearchMind AI</h1>
  <div class="doc-subtitle">Autonomous Academic Paper Deconstruction, Senior Peer Review, Grounded Chat &amp; Cross-Paper Synthesis Agent</div>
  <div class="meta-grid">
    <div class="meta-item"><strong>Platform Version</strong><span>2.0.0 (Production)</span></div>
    <div class="meta-item"><strong>AI Core Models</strong><span>Gemini 3.5 / 2.5 Flash &amp; Pro</span></div>
    <div class="meta-item"><strong>Repository</strong><span>github.com/darursrikar-sketch/researchmind-ai</span></div>
    <div class="meta-item"><strong>Document Date</strong><span>September 2026</span></div>
  </div>
</div>

<h2>1. Executive Summary &amp; Problem Statement</h2>
<p>
Modern scientific research is experiencing exponential velocity. In machine learning, computer science, and biomedicine alone, thousands of preprints are submitted to arXiv, bioRxiv, and peer-reviewed venues every week. Researchers, engineers, graduate students, and technical leaders face severe bottlenecks:
</p>
<ul>
  <li><strong>Time Constraints:</strong> Thoroughly digesting a 15–30 page dense research paper with mathematical appendices takes 2–4 hours.</li>
  <li><strong>Surface-Level LLM Summaries:</strong> Standard consumer chatbots output generic bullet points that miss crucial nuances, fail to evaluate methodological flaws, ignore missing baselines, and hallucinate unsupported assertions.</li>
  <li><strong>Neglect of Mathematical Formulations:</strong> Conventional summarizers skip LaTeX equations, loss functions, and architectural bounds entirely.</li>
  <li><strong>Fragmented Workflows:</strong> Researchers juggle multiple disconnected tools for paper retrieval, PDF annotation, chat interaction, and literature synthesis.</li>
</ul>
<div class="callout">
  <div class="callout-title">The ResearchMind Solution</div>
  <strong>ResearchMind AI</strong> is an autonomous, academic-grade AI research partner designed to bridge the gap between rapid literature scanning and rigorous peer review. Powered by Google Gemini (<code>gemini-2.5-flash</code>, <code>gemini-3.5-flash</code>, and <code>gemini-2.5-pro</code> via the official <code>google-genai</code> SDK), ResearchMind delivers multi-source ingestion, autonomous 6-pillar deep deconstruction, page-grounded conversational Q&amp;A, and cross-paper comparative meta-analysis.
</div>

<h2>2. Core Capabilities &amp; Value Proposition</h2>
<table>
  <thead>
    <tr>
      <th>Capability</th>
      <th>Standard Consumer Chatbots</th>
      <th>ResearchMind AI Platform</th>
    </tr>
  </thead>
  <tbody>
    <tr>
      <td><strong>Paper Ingestion</strong></td>
      <td>Manual copy-paste or raw file text dumps</td>
      <td>One-click arXiv ID/URL retrieval + Page-aware PDF parsing</td>
    </tr>
    <tr>
      <td><strong>Citation Precision</strong></td>
      <td>General references (prone to hallucination)</td>
      <td>Strict page-grounded citations (e.g., <code>[Page 4]</code>) verified against document text</td>
    </tr>
    <tr>
      <td><strong>Math &amp; Equations</strong></td>
      <td>Plain text mangling of Greek symbols</td>
      <td>Native KaTeX LaTeX rendering ($...$ and $$...$$) for loss functions and bounds</td>
    </tr>
    <tr>
      <td><strong>Critical Rigor</strong></td>
      <td>Polite, uncritical summaries</td>
      <td>Senior reviewer simulation (NeurIPS/ICML style, 1–10 score, rebuttal questions)</td>
    </tr>
    <tr>
      <td><strong>Replication Details</strong></td>
      <td>High-level conceptual overview</td>
      <td>Hyperparameters, compute budgets, subtle failure modes, runnable pseudocode</td>
    </tr>
    <tr>
      <td><strong>Cross-Paper Analysis</strong></td>
      <td>Limited to single context window</td>
      <td>Multi-paper comparative synthesis matrix with architectural trade-off trees</td>
    </tr>
    <tr>
      <td><strong>Streaming Delivery</strong></td>
      <td>Monolithic text chunks</td>
      <td>Real-time Server-Sent Events (SSE) streaming per analytical module</td>
    </tr>
    <tr>
      <td><strong>Export Formats</strong></td>
      <td>Raw clipboard copy</td>
      <td>Formatted GitHub Markdown, standalone print-ready HTML, structured JSON</td>
    </tr>
  </tbody>
</table>

<div class="page-break"></div>

<h2>3. System Architecture &amp; Data Flow</h2>
<p>
ResearchMind AI is structured into four decoupled, modular layers ensuring seamless scalability, robust error isolation, and maximum extensibility:
</p>

<div class="diagram-box">
+-----------------------------------------------------------------------------------+
| 1. INGESTION LAYER                                                                |
|    • arXiv API Client: Resolves Atom XML, downloads full-text PDF to local cache  |
|    • PDF Extraction Engine: Uses pypdf for page-indexed tokenization and parsing  |
+-----------------------------------------------------------------------------------+
                                         │
                                         ▼
+-----------------------------------------------------------------------------------+
| 2. AGENT CORE (Google Gemini via official google-genai SDK)                       |
|    • PaperAnalysisAgent: Orchestrates 6 structured analytical modules             |
|    • PaperChatEngine: Grounded conversational Q&amp;A enforcing exact page citations   |
|    • PaperComparisonAgent: Cross-paper comparative synthesis &amp; trade-off matrix   |
+-----------------------------------------------------------------------------------+
                                         │
                                         ▼
+-----------------------------------------------------------------------------------+
| 3. PERSISTENCE &amp; STORAGE LAYER                                                     |
|    • SQLite (data/researchmind.db): Multi-tenant papers, analyses, users, schemas |
|    • Upload Storage (data/uploads/): Retains original scientific PDFs             |
|    • Security: PBKDF2/SHA-256 salted password hashing &amp; isolated user workspaces  |
+-----------------------------------------------------------------------------------+
                                         │
                                         ▼
+-----------------------------------------------------------------------------------+
| 4. PRESENTATION &amp; CLIENT INTERFACES                                                |
|    • NexaCore Web Dashboard: Tailwind CSS, Floating Island Navbar, SSE, KaTeX     |
|    • Headless CLI: run_cli.py for batch scripts, REPL chat, and direct exports    |
|    • Exporters: Markdown (.md), Standalone HTML (.html), and JSON (.json)         |
+-----------------------------------------------------------------------------------+
</div>

<h3>3.1 Component Directory Structure</h3>
<pre><code>research mind project/
├── .env                             # Environment configuration & API credentials
├── requirements.txt                 # Python package dependencies
├── run_web.py                       # Flask web server entrypoint (Port 5000)
├── run_cli.py                       # Rich terminal CLI entrypoint
├── sample_paper.pdf                 # Pre-packaged sample scientific paper
├── sample_papers/                   # Sample benchmark library (Attention Is All You Need, etc.)
├── data/                            # Persistent data folder
│   ├── researchmind.db              # Embedded SQLite database
│   └── uploads/                     # Uploaded and fetched PDF documents
└── researchmind/                    # Primary application package
    ├── config.py                    # Environment settings & Gemini Client factory
    ├── paper_loader.py              # PDF extraction & arXiv API fetcher
    ├── storage.py                   # SQLite persistence, schema migrations & auth
    ├── exporter.py                  # Markdown, HTML, and JSON document serializers
    ├── cli.py                       # CLI command handlers & interactive REPL
    ├── core/
    │   ├── prompts.py               # Domain-specific academic prompt templates
    │   ├── agent.py                 # Multi-module PaperAnalysisAgent
    │   ├── chat_engine.py           # Grounded Q&A engine with citations
    │   └── comparative.py           # Multi-paper cross-synthesis agent
    └── web/
        ├── app.py                   # Flask REST API & SSE streaming endpoints
        ├── templates/index.html     # NexaCore-inspired reactive SPA UI
        └── static/
            ├── css/custom.css       # Design tokens, hover glows & animations
            └── js/app.js            # Client-side state, streaming parser & KaTeX hooks</code></pre>

<h2>4. The 6-Pillar Autonomous Analysis Pipeline</h2>
<p>
When an ingested paper is submitted for analysis, ResearchMind triggers an automated pipeline across six specialized analytical modules:
</p>

<div class="pillar-grid">
  <div class="pillar-card">
    <h4>1. Executive Summary &amp; TL;DR</h4>
    <p>Distills the core problem vector, novel paradigm contributions, and headline quantitative metrics into an immediate high-signal briefing.</p>
  </div>
  <div class="pillar-card">
    <h4>2. Methodology &amp; Mathematical Architecture</h4>
    <p>Step-by-step pipeline deconstruction, mathematical loss functions, tensor shapes, and theoretical proofs rendered in clean LaTeX.</p>
  </div>
  <div class="pillar-card">
    <h4>3. Critical Peer Review</h4>
    <p>Simulates a senior tier-1 reviewer (NeurIPS/ICML) with a 1–10 score, accept/reject recommendation, missing baselines, and author rebuttal questions.</p>
  </div>
  <div class="pillar-card">
    <h4>4. Findings &amp; Empirical Benchmarks</h4>
    <p>Comprehensive catalog of datasets, SotA performance comparison tables, statistical significance tests, and ablation study insights.</p>
  </div>
  <div class="pillar-card">
    <h4>5. Limitations &amp; Theoretical Bounds</h4>
    <p>Identifies fundamental failure modes, asymptotic scaling bottlenecks, memory ceilings, and unaddressed safety/adversarial concerns.</p>
  </div>
  <div class="pillar-card">
    <h4>6. Implementation &amp; Replication Blueprint</h4>
    <p>Actionable ML engineering guide featuring hyperparameter tables, GPU compute budgets, unstated implementation pitfalls, and Python pseudocode.</p>
  </div>
</div>

<div class="page-break"></div>

<h2>5. Interactive Grounded Chat &amp; Page Citations</h2>
<p>
General conversational LLMs frequently invent facts when answering technical inquiries. The <code>PaperChatEngine</code> eliminates this through strict architectural grounding:
</p>
<ul>
  <li><strong>Page-Indexed Context:</strong> The ingested paper is partitioned into page-tagged buffers. Prompts explicitly constrain responses to cited evidence.</li>
  <li><strong>Verifiable Citations:</strong> Every factual claim is accompanied by a page anchor (e.g., <code>[Page 4]</code>). The frontend renders these with distinctive badges.</li>
  <li><strong>Automated Starter Questions:</strong> Upon paper ingestion, the agent analyzes the abstract to synthesize four targeted, domain-specific exploratory questions.</li>
  <li><strong>Stateful History:</strong> Multi-turn discussions retain full context for follow-up questions while protecting against context-window saturation.</li>
</ul>

<h2>6. Cross-Paper Comparative Synthesis Matrix</h2>
<p>
The <code>PaperComparisonAgent</code> empowers researchers to select two or more papers from their local library to generate a side-by-side comparative meta-analysis:
</p>
<ol>
  <li><strong>Comparative Matrix Table:</strong> Direct juxtaposition of model architectures, objective functions, parameter scales, and benchmark scores.</li>
  <li><strong>Synergies &amp; Methodological Divergences:</strong> Identifies where papers align on foundational principles and where empirical conclusions conflict.</li>
  <li><strong>Architectural Decision Tree:</strong> Clear guidelines for practitioners indicating when to select Architecture A vs. Architecture B based on inference latency, training budget, and dataset constraints.</li>
</ol>

<h2>7. Storage &amp; Database Schema</h2>
<p>
ResearchMind stores data in an embedded SQLite database (<code>data/researchmind.db</code>) with automatic schema migration and multi-user isolation:
</p>
<table>
  <thead>
    <tr>
      <th>Table</th>
      <th>Primary Key</th>
      <th>Key Columns</th>
      <th>Description</th>
    </tr>
  </thead>
  <tbody>
    <tr>
      <td><code>users</code></td>
      <td><code>id (TEXT)</code></td>
      <td><code>username</code>, <code>email</code>, <code>password_hash</code>, <code>created_at</code></td>
      <td>User credentials for multi-tenant workspace separation. Passwords hashed using PBKDF2/SHA-256.</td>
    </tr>
    <tr>
      <td><code>papers</code></td>
      <td><code>id (TEXT)</code></td>
      <td><code>title</code>, <code>authors</code>, <code>abstract</code>, <code>arxiv_id</code>, <code>pdf_path</code>, <code>num_pages</code>, <code>data_json</code>, <code>user_id</code></td>
      <td>Ingested scientific papers, metadata, file pointers, and serialized page text mappings.</td>
    </tr>
    <tr>
      <td><code>analyses</code></td>
      <td><code>paper_id (TEXT)</code></td>
      <td><code>paper_title</code>, <code>model_used</code>, <code>data_json</code>, <code>timestamp</code>, <code>user_id</code></td>
      <td>Cached analysis outputs across the 6 pillars, preserving execution times and model identifiers.</td>
    </tr>
  </tbody>
</table>

<h2>8. REST &amp; Server-Sent Events (SSE) API Specification</h2>
<table>
  <thead>
    <tr>
      <th>Method</th>
      <th>Route</th>
      <th>Type</th>
      <th>Description</th>
    </tr>
  </thead>
  <tbody>
    <tr>
      <td><code>GET</code></td>
      <td><code>/api/auth/me</code></td>
      <td>JSON</td>
      <td>Returns active user session details or null</td>
    </tr>
    <tr>
      <td><code>POST</code></td>
      <td><code>/api/auth/signup</code></td>
      <td>JSON</td>
      <td>Registers a new user account</td>
    </tr>
    <tr>
      <td><code>POST</code></td>
      <td><code>/api/auth/signin</code></td>
      <td>JSON</td>
      <td>Authenticates credentials and establishes session cookie</td>
    </tr>
    <tr>
      <td><code>POST</code></td>
      <td><code>/api/auth/signout</code></td>
      <td>JSON</td>
      <td>Invalidates active session</td>
    </tr>
    <tr>
      <td><code>GET / POST</code></td>
      <td><code>/api/config</code></td>
      <td>JSON</td>
      <td>Reads or updates Gemini API key and default model</td>
    </tr>
    <tr>
      <td><code>GET</code></td>
      <td><code>/api/papers</code></td>
      <td>JSON</td>
      <td>Lists all papers belonging to the active workspace</td>
    </tr>
    <tr>
      <td><code>GET</code></td>
      <td><code>/api/papers/&lt;id&gt;</code></td>
      <td>JSON</td>
      <td>Retrieves paper metadata and cached analysis results</td>
    </tr>
    <tr>
      <td><code>DELETE</code></td>
      <td><code>/api/papers/&lt;id&gt;</code></td>
      <td>JSON</td>
      <td>Deletes paper and associated analyses from disk &amp; database</td>
    </tr>
    <tr>
      <td><code>POST</code></td>
      <td><code>/api/upload</code></td>
      <td>JSON</td>
      <td>Uploads and parses a multipart PDF document</td>
    </tr>
    <tr>
      <td><code>POST</code></td>
      <td><code>/api/arxiv</code></td>
      <td>JSON</td>
      <td>Fetches paper from arXiv by ID or URL</td>
    </tr>
    <tr>
      <td><code>POST</code></td>
      <td><code>/api/analyze</code></td>
      <td>SSE Stream</td>
      <td>Streams real-time analysis across selected pillars</td>
    </tr>
    <tr>
      <td><code>POST</code></td>
      <td><code>/api/chat</code></td>
      <td>SSE Stream</td>
      <td>Streams grounded conversational answers with citations</td>
    </tr>
    <tr>
      <td><code>GET</code></td>
      <td><code>/api/chat/starters/&lt;id&gt;</code></td>
      <td>JSON</td>
      <td>Generates four tailored exploratory questions</td>
    </tr>
    <tr>
      <td><code>POST</code></td>
      <td><code>/api/compare</code></td>
      <td>SSE Stream</td>
      <td>Streams comparative cross-paper synthesis matrix</td>
    </tr>
    <tr>
      <td><code>GET</code></td>
      <td><code>/api/export/&lt;id&gt;?format=md|html|json</code></td>
      <td>File / Text</td>
      <td>Downloads analysis formatted as Markdown, HTML, or JSON</td>
    </tr>
  </tbody>
</table>

<div class="page-break"></div>

<h2>9. User Interface &amp; NexaCore Design System</h2>
<p>
The web interface incorporates NexaCore's modern design language:
</p>
<ul>
  <li><strong>Floating Island Navbar (<code>.nexa-nav-island</code>):</strong> Fixed top navigation with blurred backdrop, status indicator, model selector, and gradient action buttons.</li>
  <li><strong>Dynamic Lighting Sweeps (<code>.nexa-glow-sweep</code>):</strong> Interactive cards feature smooth downward radial lighting sweeps on cursor hover.</li>
  <li><strong>Depth Physics (<code>.nexa-bottom-shade</code>):</strong> Card bases fade smoothly into dark glassmorphic shades when active.</li>
  <li><strong>Sliding Action Triggers (<code>.nexa-reveal-btn</code>):</strong> Action buttons smoothly expand from zero height on hover with vibrant gradient fills.</li>
  <li><strong>Structured Delivery Timeline:</strong> Four connected workflow milestones linked by a multi-stop vertical gradient line (<code>.nexa-grad-line-bg</code>).</li>
</ul>

<h2>10. CLI &amp; Automated Scripting Interface</h2>
<p>
For automated batch pipelines or headless Linux environments, <code>run_cli.py</code> provides a full terminal suite:
</p>
<pre><code># 1. Analyze an arXiv paper and stream findings to terminal
python run_cli.py analyze 1706.03762

# 2. Analyze a local PDF and export directly to Markdown
python run_cli.py analyze ./sample_paper.pdf --export md --out ./report.md

# 3. Analyze using Gemini 2.5 Pro for deep critical review
python run_cli.py analyze 2106.09685 --model gemini-2.5-pro --export html

# 4. Start an interactive terminal REPL chat with grounded citations
python run_cli.py chat 1706.03762

# 5. Synthesize a cross-paper comparison
python run_cli.py compare 1706.03762 2106.09685 --out comparison.md

# 6. List papers in local repository
python run_cli.py list

# 7. Configure Gemini credentials
python run_cli.py config --set-key AIzaSy... --set-model gemini-3.5-flash</code></pre>

<h2>11. Installation, Configuration &amp; Operations</h2>
<ol>
  <li><strong>Clone Repository:</strong>
    <pre><code>git clone https://github.com/darursrikar-sketch/researchmind-ai.git
cd researchmind-ai</code></pre>
  </li>
  <li><strong>Install Dependencies:</strong>
    <pre><code>pip install -r requirements.txt</code></pre>
  </li>
  <li><strong>Configure Credentials:</strong>
    <pre><code># In .env file
GEMINI_API_KEY=AIzaSy...
GEMINI_MODEL=gemini-3.5-flash
FLASK_SECRET_KEY=secure-random-key</code></pre>
  </li>
  <li><strong>Launch Application:</strong>
    <pre><code>python run_web.py</code></pre>
    Navigate to <code>http://127.0.0.1:5000</code>.
  </li>
</ol>

<h2>12. Verification &amp; Test Suite</h2>
<p>
Run the automated test suite to ensure system integrity:
</p>
<pre><code>python -m unittest discover tests</code></pre>

<div class="doc-footer">
  <strong>ResearchMind AI</strong> &bull; Empowering researchers with rigorous, grounded scientific intelligence &bull; Version 2.0.0
</div>

</body>
</html>
"""

# ==============================================================================
# 2. DOCX GENERATOR FUNCTION
# ==============================================================================

def set_cell_background(cell, hex_color):
    """Sets background color of a docx table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def create_docx_report(output_path: Path):
    doc = docx.Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)

    # Header / Title
    p_badge = doc.add_paragraph()
    r_badge = p_badge.add_run("OFFICIAL SYSTEM ARCHITECTURE & TECHNICAL REPORT")
    r_badge.font.size = Pt(9)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(79, 70, 229)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run("ResearchMind AI")
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Autonomous Academic Paper Deconstruction, Senior Peer Review, Grounded Chat & Cross-Paper Synthesis Agent")
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Metadata Table
    meta_table = doc.add_table(rows=2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("PLATFORM VERSION", "2.0.0 (Production)"),
        ("AI CORE MODELS", "Gemini 3.5 / 2.5 Flash & Pro"),
        ("REPOSITORY", "darursrikar-sketch/researchmind-ai"),
        ("DOCUMENT DATE", "September 2026")
    ]

    for idx, (label, val) in enumerate(meta_data):
        row = idx // 4
        col = idx % 4
        cell = meta_table.cell(row, col)
        set_cell_background(cell, "F1F5F9")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(f"{label}\n")
        r1.font.size = Pt(7.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(100, 116, 139)
        r2 = p.add_run(val)
        r2.font.size = Pt(9)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph() # spacer

    # Section 1
    h1 = doc.add_heading("1. Executive Summary & Problem Statement", level=1)
    h1.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Modern scientific research is experiencing exponential velocity. In machine learning, computer science, "
        "and biomedicine alone, thousands of preprints are submitted to arXiv, bioRxiv, and peer-reviewed venues every week. "
        "Researchers, engineers, graduate students, and technical leaders face severe bottlenecks:"
    )

    p1 = doc.add_paragraph(style='List Bullet')
    p1.add_run("Time Constraints: ").bold = True
    p1.add_run("Digesting a 15–30 page dense paper with mathematical appendices requires 2–4 hours.")

    p2 = doc.add_paragraph(style='List Bullet')
    p2.add_run("Surface-Level LLM Summaries: ").bold = True
    p2.add_run("Standard consumer chatbots output generic bullet points that miss nuances, fail to evaluate methodological flaws, ignore missing baselines, and hallucinate claims.")

    p3 = doc.add_paragraph(style='List Bullet')
    p3.add_run("Mathematical Neglect: ").bold = True
    p3.add_run("Conventional summarizers skip LaTeX formulas, loss functions, and architectural bounds.")

    p4 = doc.add_paragraph(style='List Bullet')
    p4.add_run("Fragmented Workflows: ").bold = True
    p4.add_run("Researchers juggle multiple disconnected tools for paper retrieval, PDF annotation, chat interaction, and literature synthesis.")

    # Callout
    p_call = doc.add_paragraph()
    p_call.paragraph_format.left_indent = Inches(0.25)
    r_call = p_call.add_run(
        "The ResearchMind AI Solution: ResearchMind AI is an autonomous, academic-grade AI research partner "
        "designed to bridge the gap between rapid literature scanning and rigorous peer review. Powered by Google Gemini "
        "(gemini-2.5-flash, gemini-3.5-flash, and gemini-2.5-pro via the official google-genai SDK), ResearchMind delivers "
        "multi-source ingestion, autonomous 6-pillar deep deconstruction, page-grounded conversational Q&A, and cross-paper comparative meta-analysis."
    )
    r_call.font.size = Pt(10)
    r_call.font.italic = True
    r_call.font.color.rgb = RGBColor(79, 70, 229)

    # Section 2
    h2 = doc.add_heading("2. Core Capabilities & Value Proposition", level=1)
    h2.style.font.color.rgb = RGBColor(15, 23, 42)

    table_comp = doc.add_table(rows=1, cols=3)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_comp.rows[0].cells
    hdr_cells[0].text = "Capability"
    hdr_cells[1].text = "Standard Consumer LLMs"
    hdr_cells[2].text = "ResearchMind AI Platform"
    for cell in hdr_cells:
        set_cell_background(cell, "0F172A")
        p = cell.paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(255, 255, 255)

    comp_data = [
        ("Paper Ingestion", "Manual copy-paste or raw file text dumps", "One-click arXiv ID/URL retrieval + Page-aware PDF parsing"),
        ("Citation Precision", "General references (prone to hallucination)", "Strict page-grounded citations ([Page 4]) verified against document text"),
        ("Math & Equations", "Plain text mangling of Greek symbols", "Native KaTeX LaTeX rendering ($...$ and $$...$$) for loss functions"),
        ("Critical Rigor", "Polite, uncritical summaries", "Senior reviewer simulation (NeurIPS/ICML style, 1–10 score, rebuttal questions)"),
        ("Replication Details", "High-level conceptual overview", "Hyperparameters, compute budgets, subtle failure modes, runnable pseudocode"),
        ("Cross-Paper Analysis", "Limited to single context window", "Multi-paper comparative synthesis matrix with architectural decision trees"),
        ("Streaming Delivery", "Monolithic text chunks", "Real-time Server-Sent Events (SSE) streaming per analytical module"),
        ("Export Formats", "Raw clipboard copy", "Formatted GitHub Markdown, standalone print-ready HTML, structured JSON")
    ]

    for row_idx, (cap, std, rm) in enumerate(comp_data):
        row_cells = table_comp.add_row().cells
        row_cells[0].text = cap
        row_cells[1].text = std
        row_cells[2].text = rm
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for cell in row_cells:
            set_cell_background(cell, bg_col)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)

    doc.add_page_break()

    # Section 3
    h3_sec = doc.add_heading("3. System Architecture & Component Design", level=1)
    h3_sec.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "ResearchMind AI is structured into four decoupled, modular layers ensuring seamless scalability, "
        "robust error isolation, and maximum extensibility:"
    )

    layers = [
        ("1. Ingestion Layer", "Includes the arXiv API Client (resolves Atom XML metadata and fetches source PDFs) and the PDF Extraction Engine (uses pypdf to maintain an ordered dictionary of page-indexed text buffers)."),
        ("2. Agent Core", "Houses the PaperAnalysisAgent (orchestrates 6 analytical modules), PaperChatEngine (stateful conversation with page citations), and PaperComparisonAgent (cross-paper meta-analysis). Directly interfaces with the official google-genai SDK."),
        ("3. Storage & Persistence", "Manages local SQLite storage (data/researchmind.db), schema migrations, PBKDF2 password hashing, isolated user workspaces, and local PDF uploads (data/uploads/)."),
        ("4. Presentation Layer", "Combines a reactive web dashboard (inspired by NexaCore with floating island navbar, radial glow sweeps, and KaTeX math), a rich terminal CLI (run_cli.py), and serialization exporters (Markdown, HTML, JSON).")
    ]

    for title, desc in layers:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{title}: ").bold = True
        p.add_run(desc)

    # Section 4
    h4_sec = doc.add_heading("4. The 6-Pillar Autonomous Analysis Pipeline", level=1)
    h4_sec.style.font.color.rgb = RGBColor(15, 23, 42)

    pillars = [
        ("Executive Summary & TL;DR", "Distills the core problem vector, novel paradigm contributions, and headline quantitative metrics into an immediate high-signal briefing."),
        ("Methodology & Mathematical Architecture", "Step-by-step pipeline deconstruction, mathematical loss functions, tensor shapes, and theoretical proofs rendered in clean LaTeX."),
        ("Critical Peer Review", "Simulates a senior tier-1 reviewer (NeurIPS/ICML) with a 1–10 score, accept/reject recommendation, missing baselines, and author rebuttal questions."),
        ("Findings & Empirical Benchmarks", "Comprehensive catalog of datasets, SotA performance comparison tables, statistical significance tests, and ablation study insights."),
        ("Limitations & Theoretical Bounds", "Identifies fundamental failure modes, asymptotic scaling bottlenecks, memory ceilings, and unaddressed safety/adversarial concerns."),
        ("Implementation & Replication Blueprint", "Actionable ML engineering guide featuring hyperparameter tables, GPU compute budgets, unstated implementation pitfalls, and Python pseudocode.")
    ]

    for title, desc in pillars:
        p = doc.add_paragraph()
        r_t = p.add_run(f"• {title}: ")
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(79, 70, 229)
        p.add_run(desc)

    # Section 5
    h5_sec = doc.add_heading("5. REST & Server-Sent Events (SSE) API", level=1)
    h5_sec.style.font.color.rgb = RGBColor(15, 23, 42)

    table_api = doc.add_table(rows=1, cols=4)
    table_api.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_hdr = table_api.rows[0].cells
    api_hdr[0].text = "Method"
    api_hdr[1].text = "Endpoint"
    api_hdr[2].text = "Format"
    api_hdr[3].text = "Function"
    for cell in api_hdr:
        set_cell_background(cell, "0F172A")
        p = cell.paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(255, 255, 255)

    api_rows = [
        ("GET", "/api/auth/me", "JSON", "Fetch active session or null"),
        ("POST", "/api/auth/signup", "JSON", "User registration"),
        ("POST", "/api/auth/signin", "JSON", "User login & session cookie"),
        ("POST", "/api/auth/signout", "JSON", "Invalidate session cookie"),
        ("GET/POST", "/api/config", "JSON", "Manage API key & default model"),
        ("GET", "/api/papers", "JSON", "List workspace papers"),
        ("GET", "/api/papers/<id>", "JSON", "Get metadata & cached analysis"),
        ("DELETE", "/api/papers/<id>", "JSON", "Delete paper & analyses"),
        ("POST", "/api/upload", "JSON", "Multipart PDF upload & extraction"),
        ("POST", "/api/arxiv", "JSON", "Ingest paper via arXiv ID/URL"),
        ("POST", "/api/analyze", "SSE Stream", "Stream real-time 6-pillar analysis"),
        ("POST", "/api/chat", "SSE Stream", "Stream grounded Q&A with citations"),
        ("GET", "/api/chat/starters/<id>", "JSON", "Generate 4 starter questions"),
        ("POST", "/api/compare", "SSE Stream", "Stream multi-paper cross synthesis"),
        ("GET", "/api/export/<id>", "File", "Export Markdown, HTML, or JSON")
    ]

    for r_idx, (m, ep, fmt, fn) in enumerate(api_rows):
        cells = table_api.add_row().cells
        cells[0].text = m
        cells[1].text = ep
        cells[2].text = fmt
        cells[3].text = fn
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for cell in cells:
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)

    # Section 6
    doc.add_page_break()
    h6_sec = doc.add_heading("6. CLI & Automated Execution Guide", level=1)
    h6_sec.style.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph("ResearchMind provides a rich command-line tool (run_cli.py) for headless and automated execution:")

    cli_examples = [
        ("Analyze arXiv Paper", "python run_cli.py analyze 1706.03762"),
        ("Analyze Local PDF & Export to Markdown", "python run_cli.py analyze ./sample_paper.pdf --export md --out ./report.md"),
        ("Deep Review with Gemini Pro", "python run_cli.py analyze 2106.09685 --model gemini-2.5-pro --export html"),
        ("Grounded Terminal Chat REPL", "python run_cli.py chat 1706.03762"),
        ("Multi-Paper Cross-Comparison", "python run_cli.py compare 1706.03762 2106.09685 --out comparison.md"),
        ("List Saved Papers in Library", "python run_cli.py list")
    ]

    for title, cmd in cli_examples:
        p = doc.add_paragraph()
        r_t = p.add_run(f"{title}:\n")
        r_t.bold = True
        r_c = p.add_run(f"  {cmd}")
        r_c.font.name = 'Consolas'
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = RGBColor(185, 28, 28)

    # Section 7
    h7_sec = doc.add_heading("7. Setup & Quick Start", level=1)
    h7_sec.style.font.color.rgb = RGBColor(15, 23, 42)

    steps = [
        "1. Install dependencies: pip install -r requirements.txt",
        "2. Configure .env with your Google Gemini API key: GEMINI_API_KEY=AIzaSy...",
        "3. Start web server: python run_web.py (Available at http://127.0.0.1:5000)",
        "4. Run unit tests: python -m unittest discover tests"
    ]
    for s in steps:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(s)

    doc.save(str(output_path))
    print(f"[SUCCESS] DOCX generated: {output_path}")

# ==============================================================================
# MAIN EXECUTION
# ==============================================================================

if __name__ == "__main__":
    html_path = BASE_DIR / "ResearchMind_Project_Documentation.html"
    docx_path = BASE_DIR / "ResearchMind_Project_Documentation.docx"
    pdf_path = BASE_DIR / "ResearchMind_Project_Documentation.pdf"

    # 1. Write HTML
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(HTML_CONTENT)
    print(f"[SUCCESS] HTML generated: {html_path}")

    # 2. Write DOCX
    create_docx_report(docx_path)

    # 3. Generate PDF via Edge Headless
    edge_exe = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
    if os.path.exists(edge_exe):
        cmd = [
            edge_exe,
            "--headless",
            "--disable-gpu",
            "--run-all-compositor-stages-before-draw",
            f"--print-to-pdf={pdf_path}",
            str(html_path.resolve()),
        ]
        try:
            res = subprocess.run(cmd, capture_output=True, timeout=30)
            if pdf_path.exists() and pdf_path.stat().st_size > 0:
                print(f"[SUCCESS] PDF generated via Edge: {pdf_path} ({pdf_path.stat().st_size} bytes)")
            else:
                print(f"[WARNING] PDF output empty or failed: {res.stderr.decode()}")
        except Exception as e:
            print(f"[ERROR] Edge print-to-pdf failed: {e}")
    else:
        print("[INFO] MS Edge not found, skipping PDF generation.")
